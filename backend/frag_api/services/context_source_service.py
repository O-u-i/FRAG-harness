from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from frag_api.schemas.context_source import (
    ExtractedSection,
    LocalFileInspectResponse,
    LocalFileSource,
    ParsedLocalFile,
    SourceLocator,
)
from frag_api.services.rag_markdown import sections_from_markdown

SUPPORTED = {".md": "md", ".txt": "txt", ".pdf": "pdf", ".docx": "docx",
             ".csv": "csv", ".xlsx": "xlsx", ".html": "html", ".htm": "html"}
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 200_000
MAX_LOCATORS = 200
CHUNK_ROWS = 50
PARSER_VERSION = "local-snapshot-v2"


class ContextSourceError(ValueError):
    pass


def _excerpt(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()[:240]


def _section(locator_id: str, kind: str, label: str, text: str) -> ExtractedSection:
    clean = text.strip()
    return ExtractedSection(
        locator=SourceLocator(id=locator_id, kind=kind, label=label, excerpt=_excerpt(clean)),
        text=clean,
    )


class ContextSourceService:
    def __init__(self, root: str | Path):
        self.root = Path(root).resolve()

    def list_files(self) -> list[LocalFileSource]:
        if not self.root.exists():
            return []
        sources: list[LocalFileSource] = []
        for path in sorted(self.root.rglob("*"), key=lambda item: item.as_posix().lower()):
            if not path.is_file() or path.name in {"README.md", ".gitkeep"}:
                continue
            try:
                resolved = self._resolve(path.relative_to(self.root).as_posix())
                sources.append(self._metadata(resolved))
            except (ContextSourceError, OSError):
                continue
        return sources

    def inspect(self, relative_path: str) -> LocalFileInspectResponse:
        parsed = self.parse(relative_path)
        return LocalFileInspectResponse(
            source=parsed.source,
            locators=[section.locator for section in parsed.sections],
            warnings=parsed.warnings,
        )

    def parse(self, relative_path: str) -> ParsedLocalFile:
        metadata, data = self.read_snapshot(relative_path)
        parser = getattr(self, f"_parse_{metadata.format}")
        sections, warnings = parser(data)
        sections = [item for item in sections if item.text]
        if len(sections) > MAX_LOCATORS:
            raise ContextSourceError("Document has too many selectable locators for direct Context.")
        if not sections:
            raise ContextSourceError("No machine-readable text was found. OCR is not supported.")
        total = sum(len(item.text) for item in sections)
        if total > MAX_EXTRACTED_CHARS:
            raise ContextSourceError("Extracted text exceeds the direct-context limit.")
        return ParsedLocalFile(source=metadata, sections=sections, warnings=warnings)

    def read_snapshot(self, relative_path: str) -> tuple[LocalFileSource, bytes]:
        path = self._resolve(relative_path)
        with path.open("rb") as handle:
            before = os.fstat(handle.fileno())
            metadata = self._metadata(path, stat=before)
            if metadata.status != "ready":
                raise ContextSourceError("File is unavailable for parsing.")
            data = handle.read(MAX_FILE_BYTES + 1)
            after = os.fstat(handle.fileno())
        fields = ("st_size", "st_mtime_ns", "st_ctime_ns", "st_ino", "st_dev")
        if len(data) > MAX_FILE_BYTES or any(getattr(before, field) != getattr(after, field) for field in fields):
            raise ContextSourceError("Source changed during snapshot capture.")
        return metadata.model_copy(update={"size_bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}), data

    def _resolve(self, relative_path: str) -> Path:
        raw = Path(relative_path)
        if raw.is_absolute() or raw.drive or raw.root or ".." in raw.parts or any(":" in part for part in raw.parts):
            raise ContextSourceError("Only relative paths inside local-sources are allowed.")
        candidate = self.root.joinpath(raw)
        current = candidate
        while current != self.root:
            if current.is_symlink() or self._is_reparse_point(current):
                raise ContextSourceError("Links are not allowed in local-sources.")
            current = current.parent
        resolved = candidate.resolve(strict=True)
        try:
            resolved.relative_to(self.root)
        except ValueError as error:
            raise ContextSourceError("Source path escapes local-sources.") from error
        current = resolved
        while current != self.root:
            if current.is_symlink() or os.path.islink(current) or self._is_reparse_point(current):
                raise ContextSourceError("Links are not allowed in local-sources.")
            current = current.parent
        if not resolved.is_file():
            raise ContextSourceError("Source must be a regular file.")
        return resolved

    @staticmethod
    def _is_reparse_point(path: Path) -> bool:
        try:
            return bool(path.lstat().st_file_attributes & 0x400)  # type: ignore[attr-defined]
        except (AttributeError, OSError):
            return False

    @staticmethod
    def _check_office_archive(data: bytes) -> None:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(item.file_size for item in archive.infolist())
            if total > 100 * 1024 * 1024:
                raise ContextSourceError("Office document expands beyond the safe parsing limit.")

    def _metadata(self, path: Path, *, stat=None) -> LocalFileSource:
        stat = stat or path.stat()
        relative = path.relative_to(self.root).as_posix()
        extension = path.suffix.lower()
        source_format = SUPPORTED.get(extension)
        status = "ready"
        reason = None
        if source_format is None:
            status, reason = "unsupported", "Unsupported file format."
        elif stat.st_size > MAX_FILE_BYTES:
            status, reason = "too_large", "File exceeds the 20 MiB parsing limit."
        return LocalFileSource(
            relativePath=relative,
            category=relative.split("/", 1)[0] if "/" in relative else "other",
            format=source_format,
            sizeBytes=stat.st_size,
            modifiedAt=datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            sha256=None,
            status=status,
            statusReason=reason,
        )

    def _parse_md(self, data: bytes):
        return [_section(item.locator_id, "section", item.title_path, item.text)
                for item in sections_from_markdown(data.decode("utf-8-sig"))], []

    def _parse_txt(self, data: bytes):
        text = data.decode("utf-8-sig")
        chunks = [chunk for chunk in re.split(r"\n\s*\n", text) if chunk.strip()]
        return [
            _section(f"section-{index + 1}", "section", f"Text block {index + 1}", chunk)
            for index, chunk in enumerate(chunks)
        ], []

    def _parse_html(self, data: bytes):
        soup = BeautifulSoup(data.decode("utf-8-sig"), "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        blocks = [tag.get_text(" ", strip=True) for tag in soup.find_all(["h1", "h2", "h3", "p", "li", "table"])]
        return [_section(f"section-{i + 1}", "section", f"HTML block {i + 1}", value) for i, value in enumerate(blocks) if value], []

    def _parse_pdf(self, data: bytes):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ContextSourceError("Encrypted PDF files are not supported.")
        if len(reader.pages) > 100:
            raise ContextSourceError("PDF exceeds the 100-page direct-context limit.")
        sections = []
        for index, page in enumerate(reader.pages):
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > 10 * 1024 * 1024:
                raise ContextSourceError("A PDF page content stream exceeds the safe parsing limit.")
            sections.append(_section(f"page-{index + 1}", "page", f"Page {index + 1}", page.extract_text() or ""))
        return sections, []

    def _parse_docx(self, data: bytes):
        self._check_office_archive(data)
        document = Document(io.BytesIO(data))
        values = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        values.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        return [_section(f"section-{i + 1}", "section", f"DOCX block {i + 1}", value) for i, value in enumerate(values)], []

    def _parse_csv(self, data: bytes):
        with io.StringIO(data.decode("utf-8-sig"), newline="") as handle:
            rows = list(csv.reader(handle))
        return self._row_sections(rows, "CSV"), []

    def _parse_xlsx(self, data: bytes):
        self._check_office_archive(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections = []
        try:
            if len(workbook.worksheets) > 10:
                raise ContextSourceError("Workbook exceeds the 10-sheet direct-context limit.")
            for sheet in workbook.worksheets:
                rows = [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
                sections.extend(self._row_sections(rows, sheet.title, prefix=f"sheet-{len(sections) + 1}"))
        finally:
            workbook.close()
        return sections, []

    @staticmethod
    def _row_sections(rows: list[list[str]], label: str, prefix: str = "rows") -> list[ExtractedSection]:
        if len(rows) > 2_000:
            raise ContextSourceError(f"{label} exceeds the 2,000-row direct-context limit.")
        if any(len(row) > 50 for row in rows):
            raise ContextSourceError(f"{label} exceeds the 50-column direct-context limit.")
        if any(len(str(cell)) > 500 for row in rows for cell in row):
            raise ContextSourceError(f"{label} contains a cell longer than 500 characters.")
        sections = []
        for start in range(0, len(rows), CHUNK_ROWS):
            batch = rows[start:start + CHUNK_ROWS]
            text = "\n".join(" | ".join(str(cell) for cell in row) for row in batch)
            sections.append(_section(f"{prefix}-{start + 1}-{start + len(batch)}", "rows", f"{label} rows {start + 1}-{start + len(batch)}", text))
        return sections
