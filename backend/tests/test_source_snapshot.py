import hashlib
from pathlib import Path

import pytest

from frag_api.services.context_source_service import ContextSourceError, ContextSourceService


def test_hash_and_extraction_use_one_byte_snapshot_even_if_path_changes(source_root, monkeypatch):
    path = source_root / "evidence.md"
    original = path.read_bytes()
    service = ContextSourceService(source_root)
    parser = service._parse_md
    def changed_path_parser(data):
        path.write_text("# Changed\nUnapproved new bytes.", encoding="utf-8")
        return parser(data)
    monkeypatch.setattr(service, "_parse_md", changed_path_parser)
    result = service.parse("evidence.md")
    assert result.source.sha256 == hashlib.sha256(original).hexdigest()
    assert "synthetic finding" in result.sections[0].text
    assert "Unapproved" not in result.sections[0].text
    assert hashlib.sha256(path.read_bytes()).hexdigest() != result.source.sha256


def test_markdown_preamble_and_fenced_heading_keep_selectable_provenance(source_root):
    text = "---\ntitle: Synthetic\n---\nPreamble fact.\n# Finding\n```text\n# Example only\n```\nReal limitation.\n## Detail\nQualified fact."
    (source_root / "evidence.md").write_text(text, encoding="utf-8")
    result = ContextSourceService(source_root).parse("evidence.md")
    assert len(result.sections) == 3
    assert result.sections[0].text == "Preamble fact."
    assert "# Example only" in result.sections[1].text and "Real limitation" in result.sections[1].text
    assert result.sections[2].locator.label == "Finding > Detail"


@pytest.mark.parametrize("path", ["../outside.txt", "C:\\outside.txt", "\\outside.txt", "/outside.txt", "evidence.md:stream", "missing.md"])
def test_snapshot_paths_reject_escape_and_alternate_stream(source_root, path):
    with pytest.raises((ContextSourceError, OSError)):
        ContextSourceService(source_root).read_snapshot(path)


def test_snapshot_formats_parse_memory_streams(source_root):
    from docx import Document
    from openpyxl import Workbook
    for name, text in {"plain.txt": "Synthetic fact.", "rows.csv": "label,value\nsynthetic,1",
                       "page.html": "<p>Synthetic fact.</p><script>password=hidden</script>"}.items():
        (source_root / name).write_text(text, encoding="utf-8")
    document = Document()
    document.add_paragraph("Synthetic document fact.")
    document.save(source_root / "document.docx")
    workbook = Workbook()
    workbook.active.append(["Synthetic", 1])
    workbook.save(source_root / "workbook.xlsx")
    workbook.close()
    service = ContextSourceService(source_root)
    for name in ("plain.txt", "rows.csv", "page.html", "document.docx", "workbook.xlsx"):
        result = service.parse(name)
        assert result.source.sha256 == hashlib.sha256((source_root / name).read_bytes()).hexdigest()
        assert result.sections
    assert "password=" not in service.parse("page.html").sections[0].text


def test_source_root_identity_cannot_be_reused_for_another_root(source_root, tmp_path):
    from frag_api.schemas.native import SourceProfile
    from frag_api.services.source_index import SourceIndex
    other = tmp_path / "other"
    other.mkdir()
    (other / "evidence.md").write_bytes((source_root / "evidence.md").read_bytes())
    first = SourceIndex(SourceProfile(id="kb",root=str(source_root),approved=True),tmp_path/"indexes")
    second = SourceIndex(SourceProfile(id="kb",root=str(other),approved=True),tmp_path/"indexes")
    a = first.scope_bindings(first.manifest())[0]
    b = second.scope_bindings(second.manifest())[0]
    assert a.revision == b.revision
    assert a.identity != b.identity
    assert first.repository.path != second.repository.path


def test_table_rows_stop_consuming_at_the_existing_limit():
    def rows():
        for _ in range(2_001):
            yield ["synthetic"]
        raise AssertionError("Oversized input must not be materialized.")
    with pytest.raises(ContextSourceError, match="2,000-row"):
        ContextSourceService._row_sections(rows(), "Synthetic")


@pytest.mark.parametrize("data,reason", [
    (b"row\n" * 2_001, "2,000-row"),
    (b",".join([b"cell"] * 51), "50-column"),
    (b"x" * 501, "500 characters"),
    ((b"x" * 500 + b"\n") * 401, "extraction limit"),
], ids=["rows", "columns", "cell-length", "total-length"])
def test_csv_limits_are_checked_during_read(source_root, data, reason):
    with pytest.raises(ContextSourceError, match=reason):
        ContextSourceService(source_root)._parse_csv(data)


def test_xlsx_oversized_dimensions_are_rejected_before_iteration(source_root, monkeypatch):
    import frag_api.services.context_source_service as parsing

    class OversizedSheet:
        max_row = 1_048_576
        max_column = 16_384

        def iter_rows(self, **_):
            raise AssertionError("Do not allocate rows for oversized dimensions.")

    class Workbook:
        worksheets = [OversizedSheet()]
        closed = False

        def close(self):
            self.closed = True

    workbook = Workbook()
    def load(_, **options):
        assert options == {"read_only": True, "data_only": True, "keep_links": False}
        return workbook
    monkeypatch.setattr(parsing, "load_workbook", load)
    monkeypatch.setattr(ContextSourceService, "_check_office_archive", lambda *_: None)
    with pytest.raises(ContextSourceError, match="dimensions"):
        ContextSourceService(source_root)._parse_xlsx(b"synthetic fixture")
    assert workbook.closed


def test_repository_requires_an_explicit_private_state_path():
    from frag_api.repositories.rag_index_repository import RagIndexRepository
    with pytest.raises(TypeError):
        RagIndexRepository()


def test_snapshot_rejects_hardlinks_to_unapproved_files(source_root, tmp_path):
    outside = tmp_path / "outside.txt"
    outside.write_text("Synthetic unapproved content.", encoding="utf-8")
    try:
        (source_root / "linked.txt").hardlink_to(outside)
    except OSError:
        pytest.skip("Hardlinks unavailable in this test environment.")
    with pytest.raises(ContextSourceError, match="regular file"):
        ContextSourceService(source_root).read_snapshot("linked.txt")


def test_git_privacy_gate_rejects_new_private_state(tmp_path):
    import subprocess
    script = Path(__file__).resolve().parents[2] / "scripts/check-repository-safety.mjs"
    subprocess.run(["git", "init", "--quiet"], cwd=tmp_path, check=True)
    result = subprocess.run(["node", str(script)], cwd=tmp_path, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    for name in [".local-guidance/plan.md", "local-sources/private-note.md", "fc.env", "nested/AGENTS.md"]:
        target = tmp_path / name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text("Synthetic privacy regression fixture.", encoding="utf-8")
        result = subprocess.run(["node", str(script)], cwd=tmp_path, capture_output=True, text=True)
        assert result.returncode != 0 and name in result.stderr
        target.unlink()
